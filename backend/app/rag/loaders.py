import os
from pathlib import Path
from typing import List
from langchain_core.documents import Document

class UnsupportedDocumentError(ValueError):
    """Raised when an unsupported document format is passed for ingestion."""
    pass

def load_pdf(file_path: str | Path) -> List[Document]:
    """
    Extracts text from a PDF document using pypdf.
    Preserves page numbers and file metadata.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PDF file not found: {file_path}")

    from pypdf import PdfReader

    documents: List[Document] = []
    try:
        reader = PdfReader(str(path))
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": str(path.resolve()),
                            "file_name": path.name,
                            "page": page_idx + 1,
                            "doc_type": "pdf",
                            "total_pages": len(reader.pages)
                        }
                    )
                )
    except Exception as e:
        raise RuntimeError(f"Error loading PDF {path.name}: {str(e)}") from e

    if not documents:
        raise ValueError(f"PDF document {path.name} contains no readable text.")

    return documents

def load_txt(file_path: str | Path) -> List[Document]:
    """
    Loads text from a plain text file.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Text file not found: {file_path}")

    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read().strip()
    except Exception as e:
        raise RuntimeError(f"Error reading TXT file {path.name}: {str(e)}") from e

    if not text:
        raise ValueError(f"Text file {path.name} is empty.")

    return [
        Document(
            page_content=text,
            metadata={
                "source": str(path.resolve()),
                "file_name": path.name,
                "page": 1,
                "doc_type": "txt"
            }
        )
    ]

def load_docx(file_path: str | Path) -> List[Document]:
    """
    Loads text from a Microsoft Word (.docx) document.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs).strip()
    except Exception as e:
        raise RuntimeError(f"Error reading DOCX file {path.name}: {str(e)}") from e

    if not text:
        raise ValueError(f"DOCX file {path.name} contains no readable text.")

    return [
        Document(
            page_content=text,
            metadata={
                "source": str(path.resolve()),
                "file_name": path.name,
                "page": 1,
                "doc_type": "docx"
            }
        )
    ]

def load_document(file_path: str | Path) -> List[Document]:
    """
    Universal document loader dispatching by extension.
    Raises UnsupportedDocumentError for unrecognized file formats.
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return load_pdf(path)
    elif ext in [".txt", ".md"]:
        return load_txt(path)
    elif ext == ".docx":
        return load_docx(path)
    else:
        supported = [".pdf", ".txt", ".md", ".docx"]
        raise UnsupportedDocumentError(
            f"Unsupported document format '{ext}' for file {path.name}. Supported: {supported}"
        )
